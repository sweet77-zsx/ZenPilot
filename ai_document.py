# -*- coding: utf-8 -*-
"""
AI 文档生成与导出模块
支持将 AI 生成的 Markdown 文本或数据汇总转换为排版精美的 Word (.docx) 及 Markdown (.md) 文档，
并提供文件存储与流式下载能力。
"""

import os
import re
import io
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generated_docs")
if not os.path.exists(DOCS_DIR):
    os.makedirs(DOCS_DIR, exist_ok=True)


def set_cell_background(cell, hex_color):
    """设置单元格背景颜色"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """设置单元格内边距 (单位为 twips, 20 twips = 1 pt)"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_cell_border(cell, **kwargs):
    """
    为单元格设置边框
    kwargs: top, bottom, left, right, etc.
    例如: set_cell_border(cell, top={"sz": 4, "val": "single", "color": "CCCCCC"})
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "CCCCCC")}"/>'
            tc_borders.append(parse_xml(tag))
    tc_pr.append(tc_borders)


def format_inline_markdown(paragraph, text, base_font_size=10.5, is_italic=False):
    """
    解析段落内的行内 Markdown（粗体 **text**、代码 `code` 等）
    """
    # 拆分 **bold** 和 `code`
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            clean_text = token[2:-2]
            run = paragraph.add_run(clean_text)
            run.bold = True
            run.font.size = Pt(base_font_size)
            run.font.name = "微软雅黑"
            run._r.get_or_add_rPr().set(qn('w:rFonts'), '微软雅黑')
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            clean_text = token[1:-1]
            run = paragraph.add_run(clean_text)
            run.font.size = Pt(base_font_size - 0.5)
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            run = paragraph.add_run(token)
            run.font.size = Pt(base_font_size)
            run.font.name = "微软雅黑"
            run._r.get_or_add_rPr().set(qn('w:rFonts'), '微软雅黑')
            if is_italic:
                run.italic = True


def markdown_to_docx(title, markdown_text, author="小禅 AI 智囊"):
    """
    将 Markdown 文本转换为排版优美的 Word (.docx) 文件流
    :param title: 文档主标题
    :param markdown_text: Markdown 格式正文
    :param author: 作者/生成人说明
    :return: BytesIO 二进制文档流
    """
    doc = docx.Document()

    # 设置页面页边距 (适度紧凑美观)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        # 页眉/页脚
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run(f"禅道桌面助手 · 小禅 AI 生成 · {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    # 1. 主标题
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run(title or "小禅 AI 智能文档")
    title_run.font.name = "微软雅黑"
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(30, 58, 138) # 深蓝

    # 2. 副标题元信息
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(16)
    meta_run = meta_p.add_run(f"汇报/生成者: {author}    |    生成日期: {datetime.now().strftime('%Y年%m月%d日 %H:%M')}")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    # 分割线
    line_p = doc.add_paragraph()
    line_p.paragraph_format.space_after = Pt(14)
    p_bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="2563EB"/></w:pBdr>')
    line_p._p.get_or_add_pPr().append(p_bdr)

    # 3. 解析 Markdown 内容
    lines = markdown_text.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        # 解析表格
        rows_data = []
        for tl in table_lines:
            # 过滤分隔线行 |---|---|
            if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', tl):
                continue
            cells = [c.strip() for c in tl.strip().strip('|').split('|')]
            if cells:
                rows_data.append(cells)

        if rows_data:
            num_rows = len(rows_data)
            num_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            for r_idx, row_vals in enumerate(rows_data):
                is_header = (r_idx == 0)
                row = table.rows[r_idx]
                for c_idx in range(num_cols):
                    cell = row.cells[c_idx]
                    val = row_vals[c_idx] if c_idx < len(row_vals) else ""
                    cell.text = "" # 清空默认段落
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    format_inline_markdown(cp, val, base_font_size=9.5 if not is_header else 10)
                    if is_header:
                        for r in cp.runs:
                            r.bold = True
                            r.font.color.rgb = RGBColor(15, 23, 42)
                        set_cell_background(cell, "E2E8F0")
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")
                    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                    set_cell_border(cell,
                        top={"sz": 4, "color": "CBD5E1"},
                        bottom={"sz": 4, "color": "CBD5E1"},
                        left={"sz": 4, "color": "CBD5E1"},
                        right={"sz": 4, "color": "CBD5E1"}
                    )

            # 表格后加点间距
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(4)
            spacer.paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        code_p = doc.add_paragraph()
        code_p.paragraph_format.space_before = Pt(4)
        code_p.paragraph_format.space_after = Pt(8)
        code_p.paragraph_format.left_indent = Inches(0.2)
        full_code = "\n".join(code_lines)
        c_run = code_p.add_run(full_code)
        c_run.font.name = "Consolas"
        c_run.font.size = Pt(9.5)
        c_run.font.color.rgb = RGBColor(30, 41, 59)
        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块起止判断
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 表格判断
        if '|' in stripped and (stripped.startswith('|') or stripped.endswith('|')):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        # 空行
        if not stripped:
            i += 1
            continue

        # 一级标题 (# Heading 1)
        if stripped.startswith("# ") and not stripped.startswith("## "):
            h_text = stripped[2:].strip()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(h_text)
            run.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # 二级标题 (## Heading 2)
        if stripped.startswith("## ") and not stripped.startswith("### "):
            h_text = stripped[3:].strip()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(h_text)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(37, 99, 235)
            i += 1
            continue

        # 三级标题 (### Heading 3)
        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(h_text)
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # 无序列表 (- or * )
        if re.match(r'^[\*\-]\s+', stripped):
            item_text = re.sub(r'^[\*\-]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_markdown(p, item_text, base_font_size=10.5)
            i += 1
            continue

        # 有序列表 (1. 2. )
        if re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_markdown(p, item_text, base_font_size=10.5)
            i += 1
            continue

        # 引用块 (> quote)
        if stripped.startswith(">"):
            quote_text = stripped.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            format_inline_markdown(p, quote_text, base_font_size=10, is_italic=True)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        format_inline_markdown(p, stripped, base_font_size=10.5)
        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def sanitize_filename(name):
    """过滤文件名非法字符"""
    clean = re.sub(r'[\\/*?:"<>|]', '_', name)
    clean = clean.strip().replace(' ', '_')
    return clean[:50] or "小禅AI文档"


def save_generated_doc(title, content, doc_type="docx", author="小禅 AI 智囊"):
    """
    生成并落盘保存文档，返回文件下载信息
    :param title: 文档标题
    :param content: Markdown 文本正文
    :param doc_type: 'docx' 或 'md'
    :param author: 作者
    :return: dict 包含文件名和下载链接
    """
    clean_title = sanitize_filename(title)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_ext = "docx" if doc_type.lower() == "docx" else "md"
    filename = f"{clean_title}_{timestamp}.{file_ext}"
    filepath = os.path.join(DOCS_DIR, filename)

    if file_ext == "docx":
        stream = markdown_to_docx(title, content, author)
        with open(filepath, "wb") as f:
            f.write(stream.getvalue())
    else:
        # Markdown 原文
        md_content = f"# {title}\n\n*作者: {author} | 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n\n---\n\n{content}"
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(md_content)

    file_size_bytes = os.path.getsize(filepath)
    size_str = f"{round(file_size_bytes / 1024, 1)} KB" if file_size_bytes >= 1024 else f"{file_size_bytes} B"

    return {
        "filename": filename,
        "filepath": filepath,
        "download_url": f"/api/ai/download-doc/{filename}",
        "title": title,
        "doc_type": file_ext,
        "file_size": size_str,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }
